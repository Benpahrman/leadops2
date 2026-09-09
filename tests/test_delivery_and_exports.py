"""Unit tests for Google Sheets diagnostics, Webhook signatures & latency, and Email CSV dispatch."""

import json
from unittest.mock import MagicMock
import pytest
from fastapi.testclient import TestClient

from agents.delivery import (
    compute_webhook_signature,
    test_webhook_connection as verify_webhook_connection,
    WebhookDestination,
    EmailCsvDestination,
)
from agents.google_sheets import (
    extract_spreadsheet_id,
    get_service_account_info,
    test_google_sheet_connection as verify_google_sheet_connection,
)
from agents.domain import Lead, State
from agents.storage import InMemoryStorageBackend
from agents.dashboard import CustomerDashboardService
from agents.api import create_app


def test_spreadsheet_id_extraction():
    url = "https://docs.google.com/spreadsheets/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/edit#gid=0"
    assert extract_spreadsheet_id(url) == "1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms"
    assert extract_spreadsheet_id("1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms") == "1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms"


def test_service_account_info():
    info = get_service_account_info()
    assert "service_account_email" in info
    assert "apps_script_template" in info
    assert "doPost" in info["apps_script_template"]


def test_google_sheets_connection_diagnostic():
    url = "https://docs.google.com/spreadsheets/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/edit"
    res = verify_google_sheet_connection(url)
    assert res["ok"] is True
    assert "spreadsheet_id" in res
    assert res["spreadsheet_id"] == "1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms"


def test_webhook_signature_calculation():
    payload = b'{"event":"test"}'
    secret = "my_secret_token_123"
    sig = compute_webhook_signature(payload, secret)
    assert isinstance(sig, str)
    assert len(sig) == 64  # SHA256 hex string


def test_webhook_destination_with_hmac():
    posted_headers = {}
    posted_body = b""

    def mock_poster(url, headers, body):
        nonlocal posted_headers, posted_body
        posted_headers = headers
        posted_body = body
        return 200

    dest = WebhookDestination(
        webhook_url="https://example.com/webhook",
        secret_token="secret_key_abc",
        http_poster=mock_poster,
    )
    rows = [{"case_number": "2026-101", "name": "Test Company"}]
    appended = dest.append(rows)
    assert appended == 1
    assert "X-LeadOps-Signature" in posted_headers
    assert "X-LeadOps-Timestamp" in posted_headers
    assert posted_headers["X-LeadOps-Secret"] == "secret_key_abc"


def test_email_csv_destination_with_mock_client():
    mock_client = MagicMock()
    dest = EmailCsvDestination(
        recipient_email="client@example.com",
        recipient_name="Client Org",
        email_sender=mock_client,
    )
    rows = [{"docket_id": "D-101", "amount": "$5,000"}]
    appended = dest.append(rows)
    assert appended == 1
    assert mock_client.send_email.called
    call_kwargs = mock_client.send_email.call_args[1]
    assert call_kwargs["to_email"] == "client@example.com"
    assert call_kwargs["is_transactional"] is True
    assert len(call_kwargs["attachments"]) == 1
    assert "leadops_feed_" in call_kwargs["attachments"][0]["filename"]


def test_api_export_and_test_destination_endpoints():
    storage = InMemoryStorageBackend()
    lead = Lead(
        lead_id="test-export-lead-99",
        tier_key="weekly",
        company_name="Test Apex Roofing",
        state=State.DELIVERED,
        contact_email="test@apexroofing.com",
    )
    storage.save_lead(lead)
    app = create_app(storage=storage)
    client = TestClient(app)

    # 1. Test Google Sheets Info endpoint
    info_resp = client.get("/api/dashboard/integrations/google-sheets-info")
    assert info_resp.status_code == 200
    assert "service_account_email" in info_resp.json()

    # 2. Test CSV Export endpoint
    csv_resp = client.get("/api/dashboard/test-export-lead-99/export")
    assert csv_resp.status_code == 200
    assert csv_resp.headers["content-type"].startswith("text/csv")

    # 3. Test JSON Export endpoint
    json_resp = client.get("/api/dashboard/test-export-lead-99/export/json")
    assert json_resp.status_code == 200
    data = json_resp.json()
    assert data["ok"] is True
    assert "records" in data

    # 4. Test Destination Ping: Google Sheets
    test_gs = client.post(
        "/api/dashboard/test-export-lead-99/test-destination",
        json={"type": "google_sheets", "url": "https://docs.google.com/spreadsheets/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/edit"},
    )
    assert test_gs.status_code == 200
    assert test_gs.json()["ok"] is True

    # 5. Test Destination Ping: Email CSV
    test_em = client.post(
        "/api/dashboard/test-export-lead-99/test-destination",
        json={"type": "email_csv", "email": "test@apexroofing.com"},
    )
    assert test_em.status_code == 200
    assert test_em.json()["ok"] is True

    # 6. Test XLSX Export endpoint
    xlsx_resp = client.get("/api/dashboard/test-export-lead-99/export/xlsx")
    assert xlsx_resp.status_code == 200
    assert "spreadsheetml" in xlsx_resp.headers["content-type"]
    assert len(xlsx_resp.content) > 1000

    # 7. Test JSONL Export endpoint
    jsonl_resp = client.get("/api/dashboard/test-export-lead-99/export/jsonl")
    assert jsonl_resp.status_code == 200
    assert "jsonlines" in jsonl_resp.headers["content-type"] or "x-ndjson" in jsonl_resp.headers["content-type"]
    lines = [line for line in jsonl_resp.text.strip().split("\n") if line.strip()]
    assert len(lines) > 0
    parsed_first_line = json.loads(lines[0])
    assert isinstance(parsed_first_line, dict)

    # 8. Test Feed Token Rotation & Public Feed Endpoints
    rotate_resp = client.post("/api/dashboard/test-export-lead-99/feed-token/rotate")
    assert rotate_resp.status_code == 200
    feed_data = rotate_resp.json()
    assert feed_data["ok"] is True
    feed_token = feed_data["feed_token"]
    assert len(feed_token) >= 16

    # 9. Test Public Feed CSV Endpoint (Zero Auth)
    public_csv = client.get(f"/api/feed/{feed_token}/records.csv")
    assert public_csv.status_code == 200
    assert public_csv.headers["content-type"].startswith("text/csv")
    assert public_csv.headers.get("access-control-allow-origin") == "*"
    assert len(public_csv.text.splitlines()) >= 2

    # 10. Test Public Feed JSON Endpoint (Zero Auth)
    public_json = client.get(f"/api/feed/{feed_token}/records.json")
    assert public_json.status_code == 200
    pj_data = public_json.json()
    assert pj_data["ok"] is True
    assert len(pj_data["records"]) > 0


def test_xlsx_and_jsonl_service_exports():
    import io
    import openpyxl
    storage = InMemoryStorageBackend()
    lead = Lead(
        lead_id="test-xlsx-export-lead",
        tier_key="weekly",
        company_name="Acme Permitting Corp",
        state=State.DELIVERED,
    )
    storage.save_lead(lead)
    svc = CustomerDashboardService(storage=storage)

    # Test XLSX Workbook Generation
    xlsx_bytes = svc.export_latest_xlsx("test-xlsx-export-lead")
    assert isinstance(xlsx_bytes, bytes)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Live Feed Data" in wb.sheetnames
    assert "Feed Metadata" in wb.sheetnames
    ws_records = wb["Live Feed Data"]
    assert ws_records.max_row >= 2
    assert ws_records.max_column >= 3

    # Test JSONL Generation
    jsonl_str = svc.export_latest_jsonl("test-xlsx-export-lead")
    assert isinstance(jsonl_str, str)
    lines = [l for l in jsonl_str.split("\n") if l.strip()]
    assert len(lines) >= 1
    for l in lines:
        obj = json.loads(l)
        assert isinstance(obj, dict)


def test_airtable_and_notion_destinations():
    from agents.delivery import (
        AirtableDestination,
        test_airtable_connection,
        NotionDestination,
        test_notion_connection,
    )
    from unittest.mock import patch, MagicMock

    # 1. Airtable Connection Test
    with patch("urllib.request.urlopen") as mock_url:
        mock_resp = MagicMock()
        mock_resp.getcode.return_value = 200
        mock_resp.read.return_value = b'{"tables": [{"name": "LeadRecords", "id": "tbl123"}]}'
        mock_url.return_value.__enter__.return_value = mock_resp

        ok, latency, msg = test_airtable_connection(
            api_key="patMockSecret123",
            base_id="appTestBase123",
            table_name="LeadRecords"
        )
        assert ok is True
        assert "LeadRecords" in msg

    # 2. Airtable Batch Append
    posted_payloads = []
    def mock_airtable_poster(url, headers, body):
        posted_payloads.append(json.loads(body.decode("utf-8")))
        return 200

    airtable_dest = AirtableDestination(
        api_key="patSecretTest",
        base_id="appTestBase",
        table_name="Leads",
        http_poster=mock_airtable_poster,
    )
    rows = [{"docket": f"DOC-{i}", "val": i * 100} for i in range(15)]
    appended = airtable_dest.append(rows)
    assert appended == 15
    # Should be 2 batches (10 + 5)
    assert len(posted_payloads) == 2
    assert len(posted_payloads[0]["records"]) == 10
    assert len(posted_payloads[1]["records"]) == 5

    # 3. Notion Connection Test
    with patch("urllib.request.urlopen") as mock_url:
        mock_resp = MagicMock()
        mock_resp.getcode.return_value = 200
        mock_resp.read.return_value = b'{"id": "db123", "title": [{"plain_text": "Live Leads"}]}'
        mock_url.return_value.__enter__.return_value = mock_resp

        ok, latency, msg = test_notion_connection(
            integration_token="secret_notion123",
            database_id="db123456789012345678901234567890"
        )
        assert ok is True
        assert "Live Leads" in msg

    # 4. Notion Page Creation
    notion_posted = []
    def mock_notion_poster(url, headers, body):
        notion_posted.append(json.loads(body.decode("utf-8")))
        return 200

    notion_dest = NotionDestination(
        api_key="secret_notion_mock",
        database_id="db456",
        http_poster=mock_notion_poster,
    )
    notion_rows = [{"Title / Case": "Case 999", "Amount": "$12,000"}]
    notion_appended = notion_dest.append(notion_rows)
    assert notion_appended == 1
    assert len(notion_posted) == 1
    assert notion_posted[0]["parent"]["database_id"] == "db456"
